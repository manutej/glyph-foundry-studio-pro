#!/usr/bin/env python3
"""Render every layout in a PowerPoint master as a labelled catalog.

The design system in a client template lives in the slide master, not the
visible slides. This makes it visible, and gives the client a way to choose
layouts by number instead of describing them in prose.

    python layout_catalog.py template.pptx [--out catalog.pptx] [--geometry]

--geometry also dumps each layout's placeholders with index, type and position
in inches, which is what you compose against.
"""
import argparse, re, sys, zipfile
from pptx import Presentation
from pptx.util import Pt, Emu


def clean_base(src, dst):
    """Strip existing slides, keep master and layouts intact."""
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                data = re.sub(rb"<p:sldIdLst>.*?</p:sldIdLst>", b"<p:sldIdLst/>",
                              data, flags=re.S)
            zout.writestr(item, data)
    print(f"  note: run the pptx skill's clean.py on {dst} to drop orphaned slide parts")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("--out", default="layout-catalog.pptx")
    ap.add_argument("--base", default="base.pptx",
                    help="where to write the slide-stripped base deck")
    ap.add_argument("--geometry", action="store_true")
    a = ap.parse_args()

    clean_base(a.template, a.base)
    prs = Presentation(a.base)

    for mi, master in enumerate(prs.slide_masters):
        lays = master.slide_layouts
        print(f"master {mi}: {len(lays)} layouts")
        for i, lay in enumerate(lays):
            s = prs.slides.add_slide(lay)
            for p in s.placeholders:
                try:
                    p.text_frame.text = f"L{i} {lay.name} | idx{p.placeholder_format.idx}"
                    for para in p.text_frame.paragraphs:
                        for r in para.runs:
                            r.font.size = Pt(11)
                except Exception:
                    pass
            if a.geometry:
                print(f'--- idx{i}  "{lay.name}"')
                for p in lay.placeholders:
                    pf = p.placeholder_format
                    print(f"    idx{pf.idx:<3} {str(pf.type):<20} "
                          f"x={Emu(p.left).inches:.2f} y={Emu(p.top).inches:.2f} "
                          f"w={Emu(p.width).inches:.2f} h={Emu(p.height).inches:.2f}")
    prs.save(a.out)
    print(f"wrote {a.out} — convert to PDF and rasterize into contact sheets to view")


if __name__ == "__main__":
    main()
